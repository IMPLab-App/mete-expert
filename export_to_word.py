from docx import Document

doc = Document()
doc.add_heading('3. Proposed Approach', level=1)
doc.add_heading('3.1. Overall Architecture of the Proposed Meta-Expert Framework', level=2)

doc.add_paragraph(
    "To tackle the inherent class imbalance in long-tailed semi-supervised learning (LTSSL), we propose a dynamic routing framework termed Meta-Expert. The core insight of our approach is to decouple the representation learning process into specialized distribution experts (e.g., a head expert and a tail expert) and dynamically aggregate their predictions using an instance-aware meta-routing mechanism. The overall architecture is built upon a standard semi-supervised pipeline (e.g., FixMatch) and mainly consists of three pivotal components: Multi-Scale Feature Extraction, Specialized Multi-Expert Classifiers, and the Meta-Expert Router."
)

doc.add_heading('1) Multi-Scale Feature Extraction Strategy', level=3)
doc.add_paragraph(
    "Given an input image x (either from the labeled set X_L or the unlabeled set X_U), we first extract its feature representations through a shared backbone network (e.g., WideResNet). Unlike conventional architectures that solely rely on the final global average pooling layer, our framework captures intermediate hierarchical representations. Formally, we denote the feature maps extracted from the k-th residual stage as f_k. These multi-scale features, denoted as F = {f_1, f_2, f_3, f_4}, encompass both low-level morphological details and high-level semantic abstractions, which are crucial for the subsequent dynamic routing phase."
)

doc.add_heading('2) Specialized Multi-Expert Classifiers', level=3)
doc.add_paragraph(
    "To mitigate the bias induced by the imbalanced label distribution, we construct multiple parallel expert branches. In our implementation, we instantiate two primary experts: a Head Expert focusing on majority classes and a Tail Expert dedicated to minority classes. Let C_h(·) and C_t(·) denote the classification heads for the head and tail experts, respectively. For the deepest feature f_4, the experts compute the corresponding logits:"
)
doc.add_paragraph("z_h = C_h(f_4), z_t = C_t(f_4)")
doc.add_paragraph(
    "To enforce expert specialization, we integrate a Logit Adjustment mechanism during the loss computation. Specifically, the predictions are calibrated using class-aware temperature scaling factors (tau_h and tau_t) based on the prior class frequency p_c:"
)
doc.add_paragraph("z_tilde_i,c = z_i,c + tau * log(p_c)")
doc.add_paragraph(
    "where tau_h is set to a standard scale (e.g., 0.0) to maintain head-class dominance, while tau_t is assigned a higher value to intrinsically penalize head classes and boost tail-class responses."
)

doc.add_heading('3) Meta-Expert Routing Mechanism', level=3)
doc.add_paragraph(
    "Relying on a static ensemble of experts is sub-optimal for instance-level variations. Therefore, we introduce a Meta-Expert router to predict sample-specific gating weights dynamically. The Meta-Expert network, parameterized by multi-layer perceptrons (MLP) with SiLU activations, takes the multi-scale feature set F and the initial expert logits as inputs. The intermediate features are embedded and concatenated to continuously maintain structural information:"
)
doc.add_paragraph("v_{meta} = MLP(Concat(E(f_1), E(f_2), E(f_3), E(f_4), z_h, z_t))")
doc.add_paragraph(
    "where E(·) represents a mapping function (e.g., linear projection). The router then outputs the adaptive gating weights w = [w_h, w_t] through a Softmax layer:"
)
doc.add_paragraph("[w_h, w_t] = Softmax(v_{meta})")
doc.add_paragraph(
    "Ultimately, the final fused discriminative output z_{fuse} for the input instance is formulated as a convex combination of the experts' logits:"
)
doc.add_paragraph("z_{fuse} = w_h * z_h + w_t * z_t")
doc.add_paragraph(
    "This formulation ensures that the network adaptively attends to the most competent expert for each specific sample during the forward propagation, significantly enhancing the model's generalization capabilities across both majority and minority classes under limited supervision."
)

doc.save('Proposed_Approach.docx')